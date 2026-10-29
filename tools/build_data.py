import openpyxl, json, collections, datetime, warnings, sys
warnings.filterwarnings('ignore')
import sys
S=sys.argv[1] if len(sys.argv)>1 else './'
if not S.endswith('/'): S+='/'
wh=openpyxl.load_workbook(S+'src/wh.xlsx',data_only=True)
_aw=openpyxl.load_workbook(S+'src/audit.xlsx',data_only=True)
au=_aw['Sheet1']
_h=list(next(au.iter_rows(min_row=4,max_row=4,values_only=True)))
AGE_LAYOUT=len(_h)>13 and _h[12]=='Cost' and _h[13]=='AGING'   # Sheet1 is the Inventory (Ageing) import
ageing=au if AGE_LAYOUT else None
pb=openpyxl.load_workbook(S+'src/pba.xlsx',data_only=True)

# warehouse code -> warehouse (as given by the business). Quarantine (-QA) and in-transit (-IT) codes are not counted.
WH_MAP={}
for _p in ("AC AC|ACEXT AC|ACEXT-MK AC|AC-MK AC|ACOW AC|ACOW-MK AC|AC-PD AC|ALAE BUKID|ARGAO ARGAO|ARGAO-PD ARGAO|BAC HOREB|"
           "BUKID BUKID|BUKID-MK BUKID|BUKID-PD BUKID|CCPC BUKID|HBEXT HOREB|HOREB HOREB|HOREB-MK HOREB|HOREB-PD HOREB|"
           "PFMIS PFMIS|PFMIS-PD PFMIS|SOUTH SOUTH|SOUTH-PD SOUTH").split('|'):
    _c,_w=_p.split(); WH_MAP[_c]=_w
def plant_of_wh(w):
    return WH_MAP.get(str(w)) if w else None
PL=['AC','PFMIS','BUKID','SOUTH','HOREB','ARGAO']
def code(x):
    try: return str(int(float(x)))
    except: return str(x)

# Warehouse of the POs raised from each PBA (RM_WH_Mgt > OPEN_PO, comment 'Origin: Blanket Agreement N.'),
# used when an OPEN PBA + PO line has no warehouse. PO_WH[(pba, item)] = {warehouse: (open qty, unit price)}
import re as _re2
PO_WH={}
for _r in wh['OPEN_PO'].iter_rows(min_row=4,values_only=True):
    _m=_re2.search(r'Blanket Agreement (\d+)',str(_r[12] or ''))
    if _m and _r[5] is not None and isinstance(_r[9],(int,float)) and _r[9]>0 and plant_of_wh(_r[7]):
        _a=PO_WH.setdefault((int(_m.group(1)),code(_r[5])),{}).setdefault(_r[7],[0.0,_r[10]]); _a[0]+=_r[9]
PO_WH={k:{w:(v[0],v[1]) for w,v in d.items()} for k,d in PO_WH.items()}
UNASSIGNED=[]
# Average cost from the Inventory Audit copy (Avg Cost column), weighted by that sheet's Ending Balance, per warehouse.
# Stock quantities stay as in RM_WH_Mgt; the audit's ending balance is only used as the weight.
# COST[(class, item code)][warehouse] = [sum(balance x avg cost), sum(balance)]
COST={}
for _r in ([] if AGE_LAYOUT else au.iter_rows(min_row=5,values_only=True)):
    if not _r[0]: continue
    _p=plant_of_wh(_r[4])
    if _p and isinstance(_r[13],(int,float)) and isinstance(_r[15],(int,float)) and _r[13]>0 and _r[15]>0:
        _a=COST.setdefault((_r[0],code(_r[1])),{}).setdefault(_p,[0.0,0.0]); _a[0]+=_r[13]*_r[15]; _a[1]+=_r[13]
# Batch cost from the Inventory (Ageing) copy: Cost (column M) weighted by QTY in KG (column J), per warehouse.
# Items without batches there have no cost (the older Inventory Audit layout is still read if Sheet1 holds that instead).
AGEC={}
if ageing is not None:
    for _r in ageing.iter_rows(min_row=5,values_only=True):
        if not _r[0]: continue
        _p=plant_of_wh(_r[4])
        if _p and isinstance(_r[9],(int,float)) and isinstance(_r[12],(int,float)) and _r[9]>0 and _r[12]>0:
            _a=AGEC.setdefault((_r[0],code(_r[3])),{}).setdefault(_p,[0.0,0.0]); _a[0]+=_r[9]*_r[12]; _a[1]+=_r[9]
def cost_for(cls,c):
    out={p:v[0]/v[1] for p,v in COST.get((cls,c),{}).items()}
    out.update({p:v[0]/v[1] for p,v in AGEC.get((cls,c),{}).items()})
    return {p:round(x,4) for p,x in out.items()}
ASOF_DT=datetime.datetime(2026,10,20)
def build(CLS,DIV):
    global items
    items={}
    def item(c,name):
        if c not in items:
            items[c]={'code':c,'name':name,'cons':dict.fromkeys(PL,0.0),'stock':dict.fromkeys(PL,0.0),'pend':dict.fromkeys(PL,0.0),'pos':[],'wh':{}}
        return items[c]

    # consumption: average weekly forecast consumption x 4, from each plant's MIN_ sheet
    mins={'MIN_BULACAN':'AC','MIN_ISA':'PFMIS','MIN_HOREB':'HOREB','MIN_ARGAO':'ARGAO','MIN_BUKID':'BUKID','MIN_SOUTH':'SOUTH','MIN_CCPC':'BUKID'}
    for sh,pl in mins.items():
        ws=wh[sh]; h=None
        for n,r in enumerate(ws.iter_rows(min_row=1,max_row=12,values_only=True),1):
            if r[0]=='INV_GROUP': h=list(r);hdr=n;break
        ix=h.index('AV W Fcasts consump')
        for r in ws.iter_rows(min_row=hdr+1,values_only=True):
            if r[0]==CLS and r[1] is not None:
                w=r[ix]
                it=item(code(r[1]),r[2]); it['cons'][pl]+=(w if isinstance(w,(int,float)) else 0)*4
    # stock (kg)
    ws=wh['INV']
    for r in ws.iter_rows(min_row=4,values_only=True):
        if r[2]==CLS and r[3] is not None:
            p=plant_of_wh(r[1])
            if p:
                it=item(code(r[3]),r[4]); it['stock'][p]+=(r[5] or 0)
                it['wh'][r[1]]=it['wh'].get(r[1],0)+(r[5] or 0)
    # pending PBA + PO
    end={}
    for r in wh['PBA'].iter_rows(min_row=4,values_only=True):
        if isinstance(r[1],(int,float)): end[(int(r[1]),code(r[6]),r[8])]=r[3]
    for r in pb['OPEN PBA + PO'].iter_rows(min_row=6,values_only=True):
        if not r[1] or r[5] is None: continue
        c=code(r[5])
        if c not in items: continue
        bal=r[9] or 0
        if bal<=0: continue
        e=end.get((int(r[1]),c,r[7]))
        # one entry per warehouse: the line's own warehouse, or (when it is blank) the warehouse of the PO raised from this PBA
        targets=[(r[7],1.0,r[11])] if plant_of_wh(r[7]) else [(w,q/sum(x[0] for x in PO_WH[(int(r[1]),c)].values()),pr) for w,(q,pr) in PO_WH[(int(r[1]),c)].items()] if (int(r[1]),c) in PO_WH else []
        if not targets or not any(plant_of_wh(t[0]) for t in targets):
            UNASSIGNED.append((CLS,c,int(r[1]),round(bal))); continue
        for w,share,pr in targets:
            p=plant_of_wh(w)
            if not p: continue
            items[c]['pend'][p]+=bal*share
            items[c]['pos'].append({'pba':int(r[1]),'vendor':r[4],'wh':w,'plant':p,'created':r[2].strftime('%Y-%m-%d'),
                'end':e.strftime('%Y-%m-%d') if e else None,'t':round(bal*share/DIV,3),'price':pr or r[11],'expired':bool(e and e<ASOF_DT)})
    out=[]
    for c,it in sorted(items.items(),key=lambda x:(0,int(x[0]),'') if x[0].isdigit() else (1,0,x[0])):
        t=lambda d:{k:round(v/DIV,3) for k,v in d.items()}
        row={'code':c,'name':it['name'],'cons':t(it['cons']),'stock':t(it['stock']),'pend':t(it['pend']),'stockWh':{k:round(v/DIV,3) for k,v in it['wh'].items() if v},
            'cost':cost_for(CLS,c),'pos':sorted(it['pos'],key=lambda x:-x['t'])}
        row['code']=c
        if sum(row['cons'].values())+sum(row['stock'].values())+sum(row['pend'].values())>0: out.append(row)
    return out

import re as _re
def _d(txt):
    m=_re.search(r'(\d{2})/(\d{2})/(\d{4})',str(txt)); return datetime.datetime(int(m.group(3)),int(m.group(1)),int(m.group(2))) if m else None
inv_sync=wh['INV']['A2'].value; pba_sync=pb['OPEN PBA + PO']['A3'].value
ASOF=_d(inv_sync) or datetime.datetime.today()
ASOF_DT=ASOF
print('inv',inv_sync,'| pba',pba_sync)
res={'asof':ASOF.strftime('%Y-%m-%d'),'invSync':ASOF.strftime('%b %d'),'invSyncFull':str(inv_sync),'costSyncFull':(str(au['A2'].value)+' | as of '+au['D2'].value.strftime('%m/%d/%Y')) if AGE_LAYOUT and hasattr(au['D2'].value,'strftime') else str(au['A2'].value),'pbaSyncFull':str(pba_sync),'pbaSync':(_d(pba_sync) or ASOF).strftime('%b %d'),'plants':PL,'groups':{
 'MACRO':{'label':'Macro','unit':'t','dec':0,'stockDec':1,'div':1000,'items':build('MACRO',1000)},
 'PACKAGING':{'label':'Packaging','unit':'pcs','dec':0,'stockDec':0,'div':1,'items':build('PACKAGING',1)},
 'MICRO':{'label':'Micro','unit':'kg','dec':0,'stockDec':0,'div':1,'items':build('MICRO',1)},
 'FEEDTAGS':{'label':'Feedtags','unit':'pcs','dec':0,'stockDec':0,'div':1,'items':build('FEEDTAGS',1)}}}
json.dump(res,open(S+'data.json','w'),separators=(',',':'))
if UNASSIGNED: print('pending lines with no warehouse found (not counted):',UNASSIGNED)
for k,g in res['groups'].items(): print(k,len(g['items']),sum(sum(i['stock'].values()) for i in g['items']),sum(sum(i['pend'].values()) for i in g['items']))
